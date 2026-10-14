"""Read the original Janesick workbook without mixing assays or annotation sheets."""

from pathlib import Path
import pandas as pd

from .cell_identity_evidence import collapse_label, normalize_label
from .epithelial_marker_review import file_hash


SHEETS = {
    'Xenium_FFPE_Human_Breast_Cancer_Rep1': 'Xenium R1 Fig1-5 (supervised)',
    'Xenium_FFPE_Human_Breast_Cancer_Rep2': 'Xenium R2 Fig1-5 (supervised)',
    'Xenium_V1_FFPE_Preview_Human_Breast_Cancer_Sample_2': 'Xenium Fig 6 (unsupervised)',
}
BROAD_NAMES = {
    'Tumor cells': 'Tumor Cells', 'Normal epithelial cells': 'Normal Epithelial Cells',
    'T cells': 'T Cells', 'B cells': 'B Cells', 'NK cells': 'T Cells',
    'Myeloid / other immune': 'pan-APCs', 'Stromal cells': 'Stromal Cells',
    'Smooth muscle cells': 'Smooth Muscle Cells', 'Other / uncertain': 'Uncertain',
}
ALIASES = {
    'trac+ cells': 'T Cells',
    'esr1+ epithelial': 'Normal Epithelial Cells',
    'pigr+ epithelial': 'Normal Epithelial Cells',
    'oprpn+ epithelial': 'Normal Epithelial Cells',
    'transitional cells': 'Uncertain',
    'not plotted': 'Uncertain', 'undefined': 'Uncertain',
}
UNLABELED = {'unlabeled', 'undefined', 'not plotted', 'unassigned', 'unknown', 'uncertain', ''}


def author_broad_label(raw):
    normalized = normalize_label(raw)
    if normalized in ALIASES:
        return ALIASES[normalized]
    broad = collapse_label(raw)[0]
    if broad not in BROAD_NAMES:
        raise ValueError(f'Unmapped Janesick author label: {raw!r}; review its mapping explicitly.')
    return BROAD_NAMES[broad]


def load_workbook(path):
    """Return each Xenium sample's original IDs, labels and explicit provenance."""
    try:
        import openpyxl
    except ImportError as error:
        raise ImportError('Reading the author workbook requires openpyxl. Install it in the notebook kernel with %pip install openpyxl.') from error

    path = Path(path)
    digest = file_hash(path)
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    tables = {}
    try:
        for sample, sheet in SHEETS.items():
            rows = book[sheet].iter_rows(values_only=True)
            header = list(next(rows))
            barcode, label = header.index('Barcode'), header.index('Cluster')
            records = []
            for row in rows:
                if all(value is None for value in row):
                    continue
                value = row[barcode]
                if value is None or isinstance(value, bool):
                    raise ValueError(f'Invalid author barcode in {sheet}')
                if isinstance(value, (int, float)):
                    if value != int(value):
                        raise ValueError(f'Noninteger author barcode in {sheet}')
                    value = str(int(value))
                raw = '' if row[label] is None else str(row[label]).strip()
                records.append((str(value).strip(), raw))
            table = pd.DataFrame(records, columns=['cell_id', 'author_raw_label'])
            if table.empty or table.cell_id.eq('').any() or table.cell_id.duplicated().any():
                raise ValueError(f'Empty sheet or blank/duplicate barcodes in {sheet}')
            table['author_refined_label'] = table.author_raw_label
            table['author_broad_label'] = table.author_raw_label.map(author_broad_label)
            table['author_label_usable'] = ~table.author_raw_label.map(normalize_label).isin(UNLABELED)
            table['author_workbook_sheet'] = sheet
            table['author_workbook_sha256'] = digest
            tables[sample] = table
    finally:
        book.close()
    if file_hash(path) != digest:
        raise ValueError('Author workbook changed while it was being read.')
    return tables


def match_author_labels(author, cell_ids):
    """Exact ID matching: never infer row offsets or borrow another sample's labels."""
    ids = pd.Index(cell_ids).astype(str)
    if ids.has_duplicates or (ids.str.strip() == '').any():
        raise ValueError('Baseline cell IDs must be unique and nonblank.')
    source = author.set_index('cell_id')
    matched = ids.isin(source.index)
    if not matched.any():
        raise ValueError('No exact workbook IDs match this sample; check input identity/ID format.')
    aligned = source.reindex(ids).copy()
    aligned.index = range(len(ids))
    aligned['author_workbook_matched'] = matched
    aligned['author_label_usable'] = aligned.author_label_usable.fillna(False).astype(bool)
    return aligned.fillna('')
